"""Verify localized delivery without changing research results or prior files."""
import os
import csv
import hashlib
import io
import json
import re
import sys
import urllib.request
from pathlib import Path
from zipfile import ZipFile

import numpy as np
import pymupdf
from docx import Document
from lxml import etree as ET

PRIVATE_IDENTITY_TOKENS = tuple(token.strip() for token in os.environ.get("PRIVATE_IDENTITY_TOKENS", "").split("|") if token.strip())

ROOT = Path(__file__).resolve().parents[3]
BUILD = ROOT / 'presentation_2026/russian/build'
SPEECH = ROOT / 'presentation_2026/speech'
version = sys.argv[1] if len(sys.argv) > 1 else 'ru-v2'
deck = BUILD / 'validated' / f'validated-{version}.pptx'
pdf = deck.with_suffix('.pdf')
source = ROOT / 'Презентация_DAMDID_Для_выступления'
out = ROOT / 'DAMDID_2026_Русский_доклад'
ns = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
sha = lambda p: hashlib.sha256(p.read_bytes()).hexdigest()
norm = lambda s: re.sub(r'\s+', '', s)
manual = json.loads((ROOT / 'revision_2026/results/manual_summary.json').read_text())
full = json.loads((ROOT / 'revision_2026/results/fullset_summary.json').read_text())
roc = json.loads((ROOT / 'presentation_2026/premium/reviews/test_roc.json').read_text())

def nums(node, kind):
    return [float(v) for v in node.xpath(f'./c:{kind}/c:numRef/c:numCache/c:pt/c:v/text() | ./c:{kind}/c:numLit/c:pt/c:v/text()', namespaces=ns)]

def near(actual, expected):
    assert len(actual) == len(expected), (actual, expected)
    assert np.allclose(actual, expected, rtol=0, atol=5.1e-9), (actual, expected)

scripts = json.loads((SPEECH / 'speech_ru.json').read_text())
with ZipFile(deck) as z, ZipFile(source / 'DAMDID_2026_Presentation.pptx') as before:
    assert z.testzip() is None
    slides = [name for name in z.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml', name)]
    notes = [name for name in z.namelist() if re.fullmatch(r'ppt/notesSlides/notesSlide\d+\.xml', name)]
    assert len(slides) == len(notes) == 20
    charts = [ET.fromstring(z.read(f'ppt/slides/charts/chart{i}.xml')) for i in range(1, 6)]
    for i, chart in enumerate(charts, 1):
        original = ET.fromstring(before.read(f'ppt/slides/charts/chart{i}.xml'))
        for path in ['//c:ser//c:numCache/c:pt/c:v/text()', '//c:ser//c:numLit/c:pt/c:v/text()']:
            near([float(v) for v in chart.xpath(path,namespaces=ns)], [float(v) for v in original.xpath(path,namespaces=ns)])
        assert chart.find('.//c:externalData', ns) is not None
    assert len([n for n in z.namelist() if n.startswith('ppt/embeddings/') and n.endswith('.xlsx')]) == 5
    for ser, label in zip(charts[0].xpath('//c:ser', namespaces=ns), ['official','human']):
        near(nums(ser,'val'), [manual['metrics'][k][label] for k in ['rrf3','bge','e5','bm25']])
    chance, curve, point = charts[1].xpath('//c:ser', namespaces=ns)
    near(nums(curve,'xVal'),[p['fpr'] for p in roc['points']])
    near(nums(curve,'yVal'),[p['tpr'] for p in roc['points']])
    depths=full['populations']['all_queries']['depths']
    for ser, key in zip(charts[2].xpath('//c:ser', namespaces=ns), ['candidates_mean_all_queries','positives_mean_all_queries']):
        near(nums(ser,'xVal'),[p['k'] for p in depths])
        near(nums(ser,'yVal'),[p[key] for p in depths])
    for chart, contrast, horizontal in zip(charts[3:],['rrf_variable_minus_e5','rrf_fixed200_minus_e5'],[True,False]):
        rows=[v for v in full['paired_inference']['all_queries'] if v['contrast']==contrast]
        mean=chart.xpath('//c:ser', namespaces=ns)[-1]
        values=nums(mean,'xVal' if horizontal else 'yVal')
        near(values,[r['mean_delta'] for r in rows])
        err=mean.find('c:errBars',ns)
        assert err.find('c:errDir',ns).get('val')==('x' if horizontal else 'y')
        for side, sign, bound in [('plus',1,1),('minus',-1,0)]:
            offsets=nums(err,side)
            near([v+sign*e for v,e in zip(values,offsets)], [r['bootstrap_ci95'][bound] for r in rows])
    text=[]
    table_owners=[]
    for i in range(1,21):
        root=ET.fromstring(z.read(f'ppt/slides/slide{i}.xml'))
        text.append(' '.join(root.xpath('//a:t/text()',namespaces=ns)))
        assert re.search('[А-Яа-яЁё]',text[-1]), i
        if root.xpath('//a:tbl',namespaces=ns): table_owners.append(i)
    assert table_owners==[2,3,4,14,16,17,18,19]
    for i,item in enumerate(scripts,1):
        note=' '.join(ET.fromstring(z.read(f'ppt/notesSlides/notesSlide{i}.xml')).xpath('//a:t/text()',namespaces=ns))
        assert all(norm(p) in norm(note) for p in item['paragraphs']),i
        assert 'Источники' in note
    for bad in ['TODO', 'ChatGPT', 'oracle-assisted', '0.597', '0.967', '0,597', '0,967', *PRIVATE_IDENTITY_TOKENS]:
        assert bad not in '\n'.join(text),bad
    assert all((name not in text[0] for name in PRIVATE_IDENTITY_TOKENS))

csv_url='https://raw.githubusercontent.com/microsoft/MSMARCO-Passage-Ranking-Submissions/abb44144f563e7972d5131f5d6acc2092877964c/leaderboard/leaderboard.csv'
csv_data=urllib.request.urlopen(csv_url).read().decode()
rows={row['id']:row for row in csv.DictReader(io.StringIO(csv_data))}
for key,value in [('20220317-SLM',.450),('20220602-fusion',.440),('20220919-cotmae',.438),('20211210-lichee-feeds',.436)]:
    assert rows[key]['type']=='full ranking'
    assert float(rows[key]['eval'])==value
    assert f'{value:.3f}'.replace('.',',') in text[2]

render_dir=BUILD/f'pdf-{version}'
render_dir.mkdir(exist_ok=True)
pdfdoc=pymupdf.open(pdf)
assert len(pdfdoc)==20
render_diffs=[]
for i,page in enumerate(pdfdoc,1):
    assert abs(page.rect.width/page.rect.height-16/9)<.0001
    pix=page.get_pixmap(matrix=pymupdf.Matrix(1.35,1.35))
    pix.save(render_dir/f'{i:02d}.png')
    previous=pymupdf.open(BUILD/'validated/validated-ru-v1.pdf')
    if pix.samples!=previous[i-1].get_pixmap(matrix=pymupdf.Matrix(1.35,1.35)).samples:
        render_diffs.append(i)
texts=[p.get_text() for p in pdfdoc]
assert '30.11.2022' in texts[2]
assert 'исторический' in texts[2]
assert all(v in texts[7] for v in ['0,248','0,488','0,433','0,459','0,644','0,879','0,828','0,854'])
assert '1,05' in texts[6] and '17,47' in texts[6]
assert '0,907' in texts[9] and '0,875' in texts[9]
assert '128,6' in texts[10] and '14,4%' in texts[12]

docx=out/'DAMDID_2026_Текст_выступления_RU.docx'
txt=out/'DAMDID_2026_Текст_выступления_RU.txt'
document=Document(docx)
assert len([p for p in document.paragraphs if p.style.name=='Heading 1'])==15
docx_text='\n'.join(p.text for p in document.paragraphs)
speech_pdf=pymupdf.open(SPEECH/'render-ru-v2/DAMDID_2026_Текст_выступления_RU.pdf')
speech_pdf_text=norm('\n'.join(p.get_text() for p in speech_pdf))
for item in scripts:
    for p in item['paragraphs']:
        assert p in docx_text and p in txt.read_text()
        assert norm(p) in speech_pdf_text, p
assert len(speech_pdf)==7
with ZipFile(docx) as z:
    for part in ['word/document.xml','word/styles.xml']:
        assert not ET.fromstring(z.read(part)).xpath('//*[local-name()="pBdr"]')

assert sha(source/'DAMDID_2026_Presentation.pptx')=='7f8e873ef396a414aeb1ef9ecad61c77a7115521dc7498b65d5e6463edaf837c'
assert sha(source/'DAMDID_2026_Presentation.pdf')=='4bbb428c036fa771e95cfddd223763644af10c884bcfaad87faa81680d335446'
report={'status':'pass','version':version,'slides':20,'main_slides':15,'appendix_slides':5,'charts_checked_against_source':5,'embedded_workbooks':5,'table_owner_slides':table_owners,'speaker_notes_synchronized':15,'speech_pages':7,'estimated_seconds':910,'timing_rehearsed':False,'pdf_pixel_differences_from_reviewed_v1':render_diffs,'leaderboard_source':csv_url,'english_source_unchanged':True,'sha256':{p.name:sha(p) for p in [deck,pdf,docx,txt]}}
(BUILD/'delivery_audit_ru.json').write_text(json.dumps(report,ensure_ascii=False,indent=2))
print(json.dumps(report,ensure_ascii=False,indent=2))
