"""Build an editable speaking script from the slide-aligned source text."""
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
BUILD = ROOT / 'presentation_2026/speech'
OUT = ROOT / 'Презентация_DAMDID_Для_выступления'
slides = json.loads((BUILD / 'speech_en.json').read_text())
assert [s['slide'] for s in slides] == list(range(1, 14))

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.85)
section.header_distance = section.footer_distance = Inches(0.32)
for name, size in [('Normal', 12), ('Title', 23), ('Subtitle', 11), ('Heading 1', 15)]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.10
    style.paragraph_format.widow_control = True
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is not None:
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn(f'w:{key}'), 'Arial')
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'en-US')
    rpr.append(language)
    for spacing in rpr.findall(qn('w:spacing')):
        rpr.remove(spacing)
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:val'), '0')
    rpr.append(spacing)
for root in (doc.styles.element, doc.element):
    for border in root.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(14)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True

body_words = sum(len(re.findall(r"\b[\w]+(?:['.-][\w]+)*\b", p)) for s in slides for p in s['paragraphs'])
durations = [round((sum(len(re.findall(r"\b[\w]+(?:['.-][\w]+)*\b", p)) for p in s['paragraphs']) * 60 / 120 + 4) / 5) * 5 for s in slides]
total = sum(durations)
def stamp(seconds):
    return f'{seconds // 60}:{seconds % 60:02d}'

doc.add_paragraph('DAMDID 2026 speaking script', 'Title')
doc.add_paragraph('Deep Re-annotation and Controlled Candidate Pools for Re-evaluating Ranking Models on MS MARCO', 'Subtitle')
intro = (
    f'English main talk for slides 1-13. Slides 14-18 are reserved for questions. '
    f'Timing is estimated at approximately 120 words per minute, with short slide transitions: '
    f'{stamp(total)} in total. Rehearse to confirm your own pace. Headings and time cues are not spoken.'
)
doc.add_paragraph(intro, 'Subtitle')
plain = ['DAMDID 2026 speaking script', '', intro, '']
timings = []
cursor = 0
for item, duration in zip(slides, durations):
    heading = f"Slide {item['slide']} {item['heading']}"
    section_heading = doc.add_paragraph(heading, 'Heading 1')
    section_heading.paragraph_format.page_break_before = item['slide'] in (4, 7, 10, 13)
    if item['slide'] in (4, 7, 10, 13):
        section_heading.paragraph_format.space_before = Pt(0)
    cue = f'{stamp(cursor)} to {stamp(cursor + duration)}'
    cue_p = doc.add_paragraph(cue, 'Subtitle')
    cue_p.paragraph_format.keep_with_next = True
    cue_p.paragraph_format.space_after = Pt(6)
    plain.extend([heading, cue, ''])
    for text in item['paragraphs']:
        doc.add_paragraph(text)
        plain.extend([text, ''])
    timings.append({'slide': item['slide'], 'start_seconds': cursor, 'duration_seconds': duration})
    cursor += duration

footer = section.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run('DAMDID 2026  |  ')
run.font.size = Pt(9)
run.font.name = 'Arial'
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'DAMDID 2026 speaking script'
doc.core_properties.subject = 'English conference talk on MS MARCO relevance judgments'
doc.core_properties.author = ''
doc.core_properties.language = 'en-US'
doc.core_properties.keywords = 'DAMDID 2026, MS MARCO, retrieval evaluation'
doc.core_properties.comments = ''
doc.core_properties.last_modified_by = ''
OUT.mkdir(exist_ok=True)
target = OUT / 'DAMDID_2026_Speech_EN.docx'
doc.save(target)
(OUT / 'DAMDID_2026_Speech_EN.txt').write_text('\n'.join(plain), encoding='utf-8')
source = ROOT / 'Презентация_DAMDID_Для_выступления/DAMDID_2026_Presentation.pdf'
report = {
    'slides': 13,
    'body_word_count': body_words,
    'estimated_seconds': total,
    'timing_assumption': '120 words/minute plus 4 seconds per slide, rounded to 5 seconds',
    'timing_is_rehearsed': False,
    'source_deck_pdf_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
    'timings': timings,
    'docx_path': str(target),
}
(BUILD / 'script_audit.json').write_text(json.dumps(report, indent=2))
print(json.dumps(report, indent=2))
