"""Build an editable speaking script from the slide-aligned source text."""
import hashlib
import json
import re
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[2]
BUILD = ROOT / 'presentation_2026/speech'
OUT = ROOT / 'DAMDID_2026_Русский_доклад'
slides = sorted(json.loads((BUILD / 'speech_ru_base.json').read_text()) + json.loads((BUILD / 'speech_intro_ru.json').read_text()), key=lambda x: x['slide'])
assert [s['slide'] for s in slides] == list(range(1, 16))
(BUILD / 'speech_ru.json').write_text(json.dumps(slides, ensure_ascii=False, indent=2))

doc = Document()
section = doc.sections[0]
section.page_width = Inches(8.27)
section.page_height = Inches(11.69)
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.7)
section.left_margin = section.right_margin = Inches(0.85)
section.header_distance = section.footer_distance = Inches(0.32)
for name, size in [('Normal', 12), ('Title', 23), ('Subtitle', 11), ('Heading 1', 15)]:
    style = doc.styles[name]
    style.font.name = 'Arial'
    style.font.size = Pt(size)
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.italic = False
    style.paragraph_format.space_after = Pt(6)
    style.paragraph_format.line_spacing = 1.10
    style.paragraph_format.widow_control = True
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is not None:
        for key in ('ascii', 'hAnsi', 'eastAsia', 'cs'):
            fonts.set(qn(f'w:{key}'), 'Arial')
    language = OxmlElement('w:lang')
    language.set(qn('w:val'), 'ru-RU')
    rpr.append(language)
    for spacing in rpr.findall(qn('w:spacing')):
        rpr.remove(spacing)
    spacing = OxmlElement('w:spacing')
    spacing.set(qn('w:val'), '0')
    rpr.append(spacing)
for root in (doc.styles.element, doc.element):
    for border in root.xpath('.//w:pBdr'):
        border.getparent().remove(border)
doc.styles['Title'].font.bold = True
doc.styles['Heading 1'].font.bold = True
doc.styles['Heading 1'].paragraph_format.space_before = Pt(14)
doc.styles['Heading 1'].paragraph_format.keep_with_next = True

body_words = sum(len(re.findall(r"\b[\w]+(?:['.-][\w]+)*\b", p)) for s in slides for p in s['paragraphs'])
durations = [round((sum(len(re.findall(r"\b[\w]+(?:['.-][\w]+)*\b", p)) for p in s['paragraphs']) * 60 / 120 + 4) / 5) * 5 for s in slides]
total = sum(durations)
def stamp(seconds):
    return f'{seconds // 60}:{seconds % 60:02d}'

doc.add_paragraph('Текст выступления DAMDID 2026', 'Title')
doc.add_paragraph('Доразметка и контролируемые пулы кандидатов для переоценки поисковых моделей на MS MARCO', 'Subtitle')
intro = (
    f'Русскоязычный доклад по слайдам 1–15. Слайды 16–20 оставлены для вопросов. '
    f'Расчётный хронометраж: {stamp(total)} при темпе около 120 слов в минуту с короткими переходами. '
    f'Точное время проверьте на репетиции. Заголовки и отметки времени не нужно зачитывать.'
)
doc.add_paragraph(intro, 'Subtitle')
plain = ['Текст выступления DAMDID 2026', '', intro, '']
timings = []
cursor = 0
for item, duration in zip(slides, durations):
    heading = f"Слайд {item['slide']} {item['heading']}"
    section_heading = doc.add_paragraph(heading, 'Heading 1')
    section_heading.paragraph_format.page_break_before = False
    cue = f'{stamp(cursor)}–{stamp(cursor + duration)}'
    cue_p = doc.add_paragraph(cue, 'Subtitle')
    cue_p.paragraph_format.keep_with_next = True
    cue_p.paragraph_format.space_after = Pt(6)
    plain.extend([heading, cue, ''])
    for index, text in enumerate(item['paragraphs']):
        paragraph = doc.add_paragraph(text)
        paragraph.paragraph_format.keep_together = True
        paragraph.paragraph_format.keep_with_next = index < len(item['paragraphs']) - 1
        plain.extend([text, ''])
    timings.append({'slide': item['slide'], 'start_seconds': cursor, 'duration_seconds': duration})
    cursor += duration

footer = section.footer.paragraphs[0]
footer.alignment = 2
run = footer.add_run('DAMDID 2026    ')
run.font.size = Pt(9)
run.font.name = 'Arial'
field = OxmlElement('w:fldSimple')
field.set(qn('w:instr'), 'PAGE')
footer._p.append(field)
doc.core_properties.title = 'Текст выступления DAMDID 2026'
doc.core_properties.subject = 'Доклад на русском языке о доразметке MS MARCO'
doc.core_properties.author = ''
doc.core_properties.language = 'ru-RU'
doc.core_properties.keywords = 'DAMDID 2026, MS MARCO, оценка поиска'
doc.core_properties.comments = ''
doc.core_properties.last_modified_by = ''
OUT.mkdir(exist_ok=True)
target = OUT / 'DAMDID_2026_Текст_выступления_RU.docx'
doc.save(target)
(OUT / 'DAMDID_2026_Текст_выступления_RU.txt').write_text('\n'.join(plain), encoding='utf-8')
source = ROOT / 'Презентация_DAMDID_Для_выступления/DAMDID_2026_Presentation.pdf'
report = {
    'slides': 15,
    'body_word_count': body_words,
    'estimated_seconds': total,
    'timing_assumption': '120 words/minute plus 4 seconds per slide, rounded to 5 seconds',
    'timing_is_rehearsed': False,
    'original_english_deck_pdf_sha256': hashlib.sha256(source.read_bytes()).hexdigest(),
    'timings': timings,
    'docx_path': str(target),
}
(BUILD / 'script_audit_ru.json').write_text(json.dumps(report, indent=2))
print(json.dumps(report, indent=2))
