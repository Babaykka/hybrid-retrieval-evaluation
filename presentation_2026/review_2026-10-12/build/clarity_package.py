import os
from pathlib import Path
import hashlib
import io
import json
import re
import shutil
import zipfile

from docx import Document
from lxml import etree as E
from PIL import Image, ImageChops
from pypdf import PdfReader

PRIVATE_IDENTITY_TOKENS = tuple(token.strip() for token in os.environ.get("PRIVATE_IDENTITY_TOKENS", "").split("|") if token.strip())


ROOT = Path(__file__).resolve().parents[3]
B = ROOT / 'presentation_2026/review_2026-10-12/build/clarity'
SOURCE = ROOT / 'DAMDID_2026_Правки_к_докладу'
OUT = ROOT / 'DAMDID_2026_Уточненная_версия'
PPT = 'DAMDID_2026_Презентация_RU.pptx'
PDF = 'DAMDID_2026_Презентация_RU.pdf'
DOC = 'DAMDID_2026_Текст_выступления_RU.docx'
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
ORDER = list(range(1, 11)) + [21, 22, 11, 23, 12, 13, 14, 15, 24, 25, 26, 16, 17, 18, 19, 20]


def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def norm(s): return ' '.join(s.split())
def read_package(path):
    with zipfile.ZipFile(path) as z:
        assert z.testzip() is None
        return {n: z.read(n) for n in z.namelist()}


parts = read_package(B / 'validated/final_v3.pptx')
old = read_package(SOURCE / PPT)
assert sha(SOURCE / PPT) == 'aab136cff068a089406bb52d54b7e2a78b01dac8e65e833c180e8b47467307d8'
assert parts.keys() == old.keys()
for name in parts:
    if '/embeddings/' in name or '/media/' in name:
        assert parts[name] == old[name], name
    if '/charts/chart' in name and name.endswith('.xml'):
        a, b = E.fromstring(parts[name]), E.fromstring(old[name])
        for expr in ('.//c:v/text()', './/c:f/text()'):
            assert a.xpath(expr, namespaces=NS) == b.xpath(expr, namespaces=NS), name

slide_text = {}
for n, part in enumerate(ORDER, 1):
    r = E.fromstring(parts[f'ppt/slides/slide{part}.xml'])
    slide_text[n] = ' '.join(r.xpath('.//a:t/text()', namespaces=NS))
assert not any('95%-ный' in s for s in slide_text.values())
assert 'ДИ 95%' in slide_text[8] and 'ДИ 95%' in slide_text[15] and 'ДИ 95%' in slide_text[22]
assert 'Псевдометки остаются' not in slide_text[9]
assert 'Судья оценивает всю dev-часть' not in slide_text[11]
assert 'Бутстрэп по запросам' not in slide_text[15]
assert 'Вывод о лучшей системе зависит от охвата разметки' in slide_text[18]
assert all((token not in str(slide_text[20]) for token in PRIVATE_IDENTITY_TOKENS))
for text in ('NDCG@10', 'RRF', 'E5'):
    assert text in slide_text[16]

doc = Document(OUT / DOC)
speeches = {}
section = None
for p in doc.paragraphs:
    if p.style.name == 'Heading 1':
        section = int(re.match(r'Слайд (\d+)', p.text)[1])
        speeches[section] = []
    elif p.style.name == 'Title': section = None
    elif section and p.style.name == 'Normal' and p.text: speeches[section].append(p.text)
for n in range(1, 21):
    notes = E.fromstring(parts[f'ppt/notesSlides/notesSlide{ORDER[n-1]}.xml'])
    body = next(sp.find('p:txBody', NS) for sp in notes.findall('.//p:sp', NS)
                if sp.find('p:nvSpPr/p:nvPr/p:ph', NS) is not None
                and sp.find('p:nvSpPr/p:nvPr/p:ph', NS).get('type') == 'body')
    text = [''.join(p.xpath('.//a:t/text()', namespaces=NS)) for p in body.findall('a:p', NS)]
    if 'Источники' in text: text = text[:text.index('Источники')]
    assert list(map(norm, text)) == list(map(norm, speeches[n])), n
text = '\n'.join(p.text for p in doc.paragraphs)
assert 'ошибки псевдоразметки' in text
assert 'queries.dev.tsv' not in text
assert 'ё' not in text and '—' not in text
assert len([p for p in doc.paragraphs if p.style.name == 'Heading 2']) == 21

unchanged = []
ref = B.parent / 'refinement/render_final'
for n in range(1, 27):
    a = Image.open(ref / f'slide-{n:02}.jpg').convert('RGB')
    b = Image.open(B / f'render_v3/slide-{n:02}.jpg').convert('RGB')
    if ImageChops.difference(a, b).getbbox() is None: unchanged.append(n)
    if n != 22:
        reviewed = Image.open(B / f'render_v2/slide-{n:02}.jpg').convert('RGB')
        assert ImageChops.difference(reviewed, b).getbbox() is None, n
assert unchanged == [1,2,3,4,5,6,7,12,13,14,17,19,20,21,23,24,25,26], unchanged
assert len(PdfReader(B / 'render_v3/final_v3.pdf').pages) == 26
assert len(PdfReader(B / f'docx_final/{Path(DOC).stem}.pdf').pages) == 12
for n in range(1, 13):
    a = Image.open(B / f'docx_render/page-{n}.png').convert('RGB')
    b = Image.open(B / f'docx_final/page-{n}.png').convert('RGB')
    assert ImageChops.difference(a, b).getbbox() is None, n

manual = json.loads((ROOT / 'revision_2026/results/manual_summary.json').read_text())
assert round(manual['metrics']['e5']['official'], 3) == 0.488
assert round(manual['metrics']['e5']['human'], 3) == 0.879
full = json.loads((ROOT / 'revision_2026/results/fullset_summary.json').read_text())
rows = full['populations']['all_queries']['depths']
deltas = {r['k']: r['models']['rrf_fixed200']['ndcg10_mean_all_queries'] - r['models']['e5']['ndcg10_mean_all_queries'] for r in rows}
assert round(deltas[10], 4) == -0.0661
assert round(deltas[30], 4) == 0.0080

shutil.copy2(B / 'validated/final_v3.pptx', OUT / PPT)
shutil.copy2(B / 'render_v3/final_v3.pdf', OUT / PDF)
forbidden = re.compile(r'__REMOVED_IDENTIFIER__|hratch|__REMOVED_IDENTIFIER__|грач|грах|бабаян', re.I)


def scan(data):
    if zipfile.is_zipfile(io.BytesIO(data)):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            assert z.testzip() is None
            for n in z.namelist():
                if not n.endswith('/'): scan(z.read(n))
    elif data.startswith(b'%PDF'):
        pdf = PdfReader(io.BytesIO(data))
        assert not forbidden.search(str(pdf.metadata))
        assert not forbidden.search('\n'.join(p.extract_text() or '' for p in pdf.pages))
    else:
        for encoding in ('utf-8', 'utf-16-le', 'utf-16-be'):
            try: text = data.decode(encoding)
            except UnicodeError: continue
            assert not forbidden.search(text)


for name in (PPT, PDF, DOC): scan((OUT / name).read_bytes())
archive = OUT.with_suffix('.zip')
with zipfile.ZipFile(archive, 'x', zipfile.ZIP_DEFLATED) as z:
    for name in (PPT, PDF, DOC): z.write(OUT / name, name)
with zipfile.ZipFile(archive) as z:
    assert z.testzip() is None
    assert len(z.namelist()) == 3
    for name in (PPT, PDF, DOC): assert z.read(name) == (OUT / name).read_bytes()
scan(archive.read_bytes())
report = {'status': 'pass', 'pixel_identical_slides': unchanged,
          'changed_slides_visually_reviewed': [8,9,10,11,15,16,18,22],
          'notes_match_speech': True, 'docx_pages_reviewed': 12,
          'metadata_scan': 'pass', 'numbers_and_workbooks_preserved': True,
          'files': {n: sha(OUT / n) for n in (PPT, PDF, DOC)},
          'archive': str(archive), 'archive_sha256': sha(archive)}
(B / 'delivery_verified.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps(report, ensure_ascii=False, indent=2))
