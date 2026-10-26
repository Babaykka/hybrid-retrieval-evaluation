import os
from pathlib import Path
from collections import defaultdict
import hashlib
import io
import json
import re
import shutil
import zipfile

from docx import Document
from lxml import etree as E
from PIL import Image, ImageChops
from pypdf import PdfReader

PRIVATE_IDENTITY_TOKENS = tuple(token.strip() for token in os.environ.get("PRIVATE_IDENTITY_TOKENS", "").split("|") if token.strip())

ROOT = Path(__file__).resolve().parents[3]
B = ROOT / 'presentation_2026/review_2026-10-12/build/refinement'
OUT = ROOT / 'DAMDID_2026_Правки_к_докладу'
PPT = 'DAMDID_2026_Презентация_RU.pptx'
PDF = 'DAMDID_2026_Презентация_RU.pdf'
DOC = 'DAMDID_2026_Текст_выступления_RU.docx'
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
ORDER = list(range(1, 11)) + [21, 22, 11, 23, 12, 13, 14, 15, 24, 25, 26, 16, 17, 18, 19, 20]


def norm(s): return ' '.join(s.split())
def sha(p): return hashlib.sha256(p.read_bytes()).hexdigest()


with zipfile.ZipFile(B / 'validated/final_v3.pptx') as z:
    parts = {n: z.read(n) for n in z.namelist()}
    assert z.testzip() is None
with zipfile.ZipFile(ROOT / 'DAMDID_2026_После_рецензии' / PPT) as z:
    old = {n: z.read(n) for n in z.namelist()}
assert parts.keys() == old.keys()
for name in parts:
    if '/embeddings/' in name or '/media/' in name:
        assert parts[name] == old[name], name
    if '/charts/chart' in name and name.endswith('.xml'):
        a, b = E.fromstring(parts[name]), E.fromstring(old[name])
        for expr in ('.//c:v/text()', './/c:f/text()'):
            assert a.xpath(expr, namespaces=NS) == b.xpath(expr, namespaces=NS)

doc = Document(OUT / DOC)
speeches = {}
section = None
for p in doc.paragraphs:
    if p.style.name == 'Heading 1':
        section = int(re.match(r'Слайд (\d+)', p.text)[1])
        speeches[section] = []
    elif p.style.name == 'Title': section = None
    elif section and p.style.name == 'Normal' and p.text: speeches[section].append(p.text)
for n in range(1, 21):
    notes = E.fromstring(parts[f'ppt/notesSlides/notesSlide{ORDER[n-1]}.xml'])
    body = next(sp.find('p:txBody', NS) for sp in notes.findall('.//p:sp', NS)
                if sp.find('p:nvSpPr/p:nvPr/p:ph', NS) is not None
                and sp.find('p:nvSpPr/p:nvPr/p:ph', NS).get('type') == 'body')
    text = [''.join(p.xpath('.//a:t/text()', namespaces=NS)) for p in body.findall('a:p', NS)]
    text = text[:text.index('Источники')]
    assert list(map(norm, text)) == list(map(norm, speeches[n])), n
assert 'Оранжевые столбцы' in ' '.join(speeches[8])
assert 'Серые столбцы' not in ' '.join(speeches[8])
assert 'queries.dev.tsv' not in '\n'.join(p.text for p in doc.paragraphs)
assert all((token not in str(' '.join(speeches[20])) for token in PRIVATE_IDENTITY_TOKENS))

manual = json.loads((ROOT / 'revision_2026/results/manual_summary.json').read_text())
bad = {r[0] for r in manual['invalid_labels']}
empty = {r[0] for r in manual['empty_texts']}
assert (len(manual['invalid_labels']), len(bad), len(empty), len(bad | empty)) == (29, 18, 5, 23)
full = json.loads((ROOT / 'revision_2026/results/fullset_summary.json').read_text())
all_rows = full['populations']['all_queries']['depths']
excluded_rows = full['populations']['excluding_manual_queries']['depths']
deltas = []
for a, b in zip(all_rows, excluded_rows):
    for model in ('rrf_variable', 'rrf_fixed200'):
        values = [r['models'][model]['ndcg10_mean_all_queries'] - r['models']['e5']['ndcg10_mean_all_queries'] for r in (a, b)]
        assert values[0] * values[1] > 0
        deltas.append({'K': a['k'], 'model': model, 'before': values[0], 'after': values[1]})
assert full['populations']['excluding_manual_queries']['query_count'] == 101093 - 1179
rrf = all_rows[0]['models']['rrf_fixed200']
assert round(100 * rrf['top10_outside_pool_total'] / rrf['top10_count_total'], 1) == 14.4
rankings = [list(range(i * 10, (i + 1) * 10)) + [999] for i in range(3)]
def fuse(depth):
    scores = defaultdict(float)
    for ranking in rankings:
        for rank, docid in enumerate(ranking[:depth], 1): scores[docid] += 1 / (60 + rank)
    return sorted(scores, key=lambda d: (-scores[d], d))[:10]
assert 999 not in fuse(10) and fuse(11)[0] == 999

unchanged = []
for n in range(1, 27):
    ref = Image.open(B / f'reference_render/slide-{n:02}.jpg').convert('RGB')
    image = Image.open(B / f'render_final/slide-{n:02}.jpg').convert('RGB')
    if ImageChops.difference(ref, image).getbbox() is None: unchanged.append(n)
assert unchanged == [1,2,3,4,5,6,7,9,10,12,13,19,20,21,22,23,24]
assert len(PdfReader(B / 'render_final/final_v3.pdf').pages) == 26
assert len(PdfReader(B / 'docx_final/DAMDID_2026_Текст_выступления_RU.pdf').pages) == 11

shutil.copy2(B / 'validated/final_v3.pptx', OUT / PPT)
shutil.copy2(B / 'render_final/final_v3.pdf', OUT / PDF)
forbidden = re.compile(r'__REMOVED_IDENTIFIER__|hratch|__REMOVED_IDENTIFIER__|грач|грах|бабаян', re.I)
def scan(data):
    if zipfile.is_zipfile(io.BytesIO(data)):
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            assert z.testzip() is None
            for n in z.namelist():
                if not n.endswith('/'): scan(z.read(n))
    elif data.startswith(b'%PDF'):
        pdf = PdfReader(io.BytesIO(data))
        assert not forbidden.search(str(pdf.metadata))
        assert not forbidden.search('\n'.join(p.extract_text() or '' for p in pdf.pages))
    else:
        for encoding in ('utf-8','utf-16-le','utf-16-be'):
            try: text = data.decode(encoding)
            except UnicodeError: continue
            assert not forbidden.search(text)
for name in (PPT, PDF, DOC): scan((OUT / name).read_bytes())
archive = OUT.with_suffix('.zip')
with zipfile.ZipFile(archive, 'x', zipfile.ZIP_DEFLATED) as z:
    for name in (PPT, PDF, DOC): z.write(OUT / name, name)
with zipfile.ZipFile(archive) as z:
    assert z.testzip() is None
    for name in (PPT, PDF, DOC): assert z.read(name) == (OUT / name).read_bytes()
scan(archive.read_bytes())
report = {'status': 'pass', 'pixel_identical_slides': unchanged, 'changes_visually_reviewed': [8,11,14,15,16,17,18,25,26],
          'docx_pages_reviewed': 11, 'notes_match_speech': True, 'rrf_example_verified': True,
          'exclusion_signs_verified': deltas, 'metadata_scan': 'pass',
          'files': {name: sha(OUT / name) for name in (PPT, PDF, DOC)}, 'archive': str(archive), 'archive_sha256': sha(archive)}
(B / 'delivery_verified.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps({'status': report['status'], 'archive': str(archive), 'files': list(report['files']), 'unchanged_slide_count': len(unchanged)}, ensure_ascii=False))
