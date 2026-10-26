from pathlib import Path
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
import math
import re
import zipfile

from docx import Document
from docx.shared import Pt
from lxml import etree as E


ROOT = Path(__file__).resolve().parents[3]
BUILD = ROOT / 'presentation_2026/review_2026-10-12/build/refinement'
SOURCE = ROOT / 'DAMDID_2026_После_рецензии'
OUT = ROOT / 'DAMDID_2026_Правки_к_докладу'
PPT = 'DAMDID_2026_Презентация_RU.pptx'
DOC = 'DAMDID_2026_Текст_выступления_RU.docx'
BUILD.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)
assert hashlib.sha256((SOURCE / PPT).read_bytes()).hexdigest() == 'a9fbc8930dbf18c31e09149dc9ef9a17c27bc3a4ef231cc8fdac2c6a6d5f6bca'

NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
ORDER = list(range(1, 11)) + [21, 22, 11, 23, 12, 13, 14, 15, 24, 25, 26, 16, 17, 18, 19, 20]


def clean(text):
    text = text.replace('ё', 'е').replace('Ё', 'Е').replace('—', '-')
    return re.sub(r'(?<=\d) (?=\d{3}(?:\D|$))', '\u00a0', text)


source_doc = Document(SOURCE / DOC)
speech, titles, qa = {}, {}, []
section, question = None, None
for p in source_doc.paragraphs:
    if p.style.name == 'Heading 1':
        match = re.match(r'Слайд (\d+)\s+(.+)', p.text)
        section = int(match[1])
        speech[section] = []
        titles[section] = clean(match[2])
    elif p.style.name == 'Heading 2':
        section = None
        question = {'q': clean(p.text), 'answer': [], 'slide': None}
        qa.append(question)
    elif question is not None and p.style.name == 'Subtitle':
        question['slide'] = int(re.search(r'\d+', p.text)[0])
    elif question is not None and p.style.name == 'Normal' and p.text:
        question['answer'].append(clean(p.text))
    elif section is not None and p.style.name == 'Normal' and p.text:
        speech[section].append(clean(p.text))
speech[20] = ['Работа выполнена при поддержке Российского научного фонда по соглашению № __REMOVED_IDENTIFIER__. Спасибо за внимание! Буду рад вашим вопросам.']
speech[8] = [p.replace('Серые столбцы', 'Оранжевые столбцы') for p in speech[8]]
speech[11] = [
    'Теперь переходим от ручной подвыборки к полной dev-части MS MARCO Passage Ranking. Она содержит 101 093 запроса. Для каждого запроса поисковые модели заранее нашли фрагменты-кандидаты, а обученный судья оценивает их релевантность. Это число запросов; фрагментов и пар «запрос-фрагмент» намного больше.',
    'Внутри полной dev-части есть стандартная оценочная подвыборка dev/small из 6 980 запросов. Из нее взяты наши 1 179 запросов с расширенной ручной разметкой. Судью дообучали только на обучающей части этой ручной подвыборки. Теперь используем его предсказания для кандидатов по всем 101 093 запросам. Человеческими метками эти предсказания не подменяем. Отдельно проверим, сохраняются ли результаты без запросов из ручной подвыборки.'
]
speech[14] = [speech[14][0],
    'RRF означает слияние по обратным рангам. Для каждого фрагмента смотрим его место в выдаче BM25, E5 и BGE. Каждый список, в котором он найден, добавляет один, деленный на шестьдесят плюс место. Эти вклады складываем и сортируем фрагменты по сумме. У трех систем одинаковый вес.',
    'Например, фрагмент попал в первые L результатов E5 и BGE, но не попал в первые L результатов BM25. Тогда складываются только два вклада: от E5 и BGE. BM25 этому фрагменту ничего не добавляет. Фрагмент при этом остается кандидатом для итоговой выдачи RRF.',
    'Поэтому RRF объединяет позиции, а не усредняет исходные баллы трех моделей. Это удобно, потому что их баллы имеют разные шкалы. На выходе получаем один итоговый ранжированный список. Оценки судьи не участвуют в его построении: они нужны для последующего сравнения качества.'
]
speech[15].insert(1,
    'Почему берем больше десяти результатов, хотя считаем NDCG по первой десятке? Потому что десять в NDCG относится к итоговому списку после слияния. Фрагмент, стоящий одиннадцатым у всех трех систем, при L, равном десяти, вообще не попадет в RRF. При L не меньше одиннадцати он получит сразу три вклада и может войти в итоговый топ-10. Поэтому более глубокие входные списки способны изменить и состав первой десятки, и порядок внутри нее.')
speech[16][2] = (
    'При K, равном десяти, 14,4 процента пар из итогового топ-10 RRF лежат за пределами пула с оценками. При расчете метрики они дают нулевой вклад. Но отсутствие оценки не означает, что фрагмент нерелевантен: в этом пуле мы его просто не оценивали. Разметка более глубокого пула позволяет проверить и такие фрагменты.')
speech[17] = [
    'В масштабном эксперименте все метки получены одним судьей. Это относится и к тем запросам, которые раньше размечал человек. Ручные ответы здесь не заменяют предсказания судьи и не дополняют их.',
    'Однако часть этих запросов использовалась при дообучении судьи. Поэтому мы дополнительно пересчитали сравнение, исключив все 1 179 запросов ручной подвыборки. Из 101 093 осталось 99 914. На каждой проверенной глубине преимущество осталось у той же системы, что и до исключения. Это говорит об устойчивости сравнения к такому пересечению, но не заменяет независимый тест судьи.',
    'Ограничения сохраняются. Человек видел известный релевантный ответ, а судья получает только запрос и кандидата. Согласие между разметчиками не измерялось. Более глубокие пулы не проходили независимую ручную проверку.',
    'Наши численные выводы относятся к выбранным англоязычным пулам MS MARCO. Их перенос на другие языки, предметные области и поисковые сервисы еще предстоит проверить.'
]
speech[18][2] = (
    'Поэтому вместе с метрикой качества поиска нужно указывать пул разметки, источник меток и настройки слияния. Направления дальнейшего исследования: сравнить более сильных судей, проверить пользу известного релевантного ответа на входе и оценить результаты по независимой человеческой разметке. Результатов этих новых экспериментов пока нет.')

qa[0]['answer'] = ['101 093 запроса содержит полная dev-часть MS MARCO Passage Ranking v1. Мы использовали именно этот набор, без отбора до 101 тысячи из более крупной выборки. Его стандартная оценочная подвыборка dev/small содержит 6 980 запросов. Число около 140 тысяч к этому набору не относится.']
qa[10]['answer'] = ['Для E5 и BGE сохранены выдачи для всех 101 093 запросов. У BM25 нет выдач для трех запросов, а некоторые списки короче заданной глубины. Используем доступные результаты. Если фрагмент не попал в первые L результатов одной системы, она не добавляет ему балл RRF, но другие системы могут его добавить.']
qa[13] = {
    'q': 'В масштабном эксперименте смешаны оценки человека и судьи?',
    'slide': 17,
    'answer': [
        'Нет. Все масштабные метрики рассчитаны по псевдометкам одного судьи. Это касается и 1 179 запросов, для которых есть ручная разметка. В этой части расчета ручные метки не подставлялись вместо предсказаний и не объединялись с ними. Сравнение исходных, человеческих и машинных меток относится к отдельному ручному эксперименту.',
        'Полная dev-часть включает запросы, участвовавшие в дообучении судьи. Поэтому дополнительно исключили все 1 179 ручных запросов, включая обучающие, валидационные и тестовые, и пересчитали метрики на оставшихся 99 914. Для каждого K знак разности RRF и E5 сохранился. Это проверка устойчивости сравнения. Все 101 093 запроса нельзя называть независимым тестом судьи.'
    ]
}
qa[15]['slide'] = 17
qa.extend([
    {'q': 'Зачем L больше 10, если метрика NDCG@10?', 'slide': 15,
     'answer': [
         'L ограничивает входные списки BM25, E5 и BGE до слияния. NDCG@10 оценивает первые десять позиций уже получившегося списка RRF. Это разные этапы.',
         'Условный пример: в первых десяти результатах трех систем нет общих фрагментов, но один и тот же фрагмент стоит одиннадцатым у всех трех. При L = 10 его нет среди кандидатов RRF. При L = 11 его сумма равна 3/(60 + 11), примерно 0,0423. У фрагмента, который стоит первым только в одном списке, сумма 1/(60 + 1), примерно 0,0164. Общий фрагмент становится первым в RRF. Поэтому L > 10 может менять итоговый топ-10. При этом дальнейшее увеличение L не обязано менять выдачу на каждом запросе.'
     ]},
    {'q': 'Что именно не так с 23 запросами ручной подвыборки?', 'slide': 25,
     'answer': [
         'В исходных записях нашли 29 пар с кодами меток вне принятой шкалы 0/1/2: 26 меток со значением 3, две со значением 20 и одну со значением 22. Эти записи относятся к 18 запросам. Значение таких кодов не установлено, поэтому их нельзя произвольно считать подтвержденными положительными или отрицательными оценками.',
         'Еще у пяти других запросов обнаружили по одному кандидату с пустым текстом. Всего затронуты 23 запроса. В дополнительном расчете исключили эти запросы целиком, осталось 1 156. Порядок поисковых систем по человеческим оценкам не изменился. В основном расчете нестандартные метки не считались подтвержденными положительными оценками и давали нулевой вклад. Исключение служит проверкой чувствительности результата, а не новой разметкой.'
     ]},
    {'q': 'Неоцененный фрагмент считается нерелевантным?', 'slide': 16,
     'answer': ['При расчете NDCG фрагмент без оценки в текущем пуле дает нулевой вклад. Это правило подсчета при неполной разметке, а не установленная нерелевантность. Часть таких фрагментов может оказаться полезной после оценки. При K = 10 за пределами пула лежат 14,4% пар из топ-10 RRF, построенного из списков глубины L = 200.']}
])

with zipfile.ZipFile(SOURCE / PPT) as z:
    parts = {n: z.read(n) for n in z.namelist()}
original = dict(parts)


def tree(part):
    return E.fromstring(parts[f'ppt/slides/slide{part}.xml'])


def xml_bytes(node):
    return E.tostring(node, encoding='UTF-8', xml_declaration=True, standalone=True)


def text_body(body, paragraphs):
    old = body.findall('a:p', NS)
    template = old[0] if old else None
    for p in old:
        body.remove(p)
    for line in paragraphs:
        p = E.SubElement(body, f'{{{NS["a"]}}}p')
        if template is not None and template.find('a:pPr', NS) is not None:
            p.append(deepcopy(template.find('a:pPr', NS)))
        r = E.SubElement(p, f'{{{NS["a"]}}}r')
        if template is not None and template.find('a:r/a:rPr', NS) is not None:
            r.append(deepcopy(template.find('a:r/a:rPr', NS)))
        E.SubElement(r, f'{{{NS["a"]}}}t').text = clean(line)


def shape_text(root, shape_id, lines):
    sp = next(s for s in root.findall('.//p:sp', NS)
              if s.find('p:nvSpPr/p:cNvPr', NS).get('id') == str(shape_id))
    text_body(sp.find('p:txBody', NS), lines.split('\n'))


changes = {
    21: {4: 'запроса в полной dev-части',
         8: 'Судья оценивает всю dev-часть; дообучение проходило на ручной подвыборке.'},
    23: {14: 'Равный вес источников. Учитываем списки, в которых найден фрагмент.'},
    12: {4: 'L = K. Более глубокие списки могут менять итоговый топ-10.'},
    13: {11: 'Отсутствие оценки дает нулевой вклад в метрику, но не означает нерелевантность.'},
    14: {3: 'В обоих расчетах используются только псевдометки судьи'},
    15: {10: 'Направления дальнейшего исследования\nБолее сильные судьи, эталонный ответ, независимая ручная проверка.'},
    19: {8: 'Проверка качества ручных данных',
         3: 'Ручная разметка: проверка без 23 запросов',
         5: '29 пар с метками 3, 20 или 22; 5 пар без текста',
         6: 'Затронуты 23 запроса. После исключения осталось 1 156.',
         7: 'Коды вне шкалы 0/1/2 не считали подтвержденными оценками.'},
    20: {9: 'Судья с эталонным ответом'},
}
for part, edits in changes.items():
    root = tree(part)
    for shape_id, text in edits.items():
        shape_text(root, shape_id, text)
    if part == 14:
        table = root.find('.//a:tbl', NS)
        rows = table.findall('a:tr', NS)
        values = [
            ['Запросы', 'Число', 'Сравнение RRF и E5'],
            ['Полная dev-часть', '101 093', 'Базовое сравнение'],
            ['Без 1 179 запросов\nручной подвыборки', '99 914', 'На тех же K лидирует\nта же система']
        ]
        for row, row_values in zip(rows, values):
            for cell, value in zip(row.findall('a:tc', NS), row_values):
                text_body(cell.find('a:txBody', NS), value.split('\n'))
    parts[f'ppt/slides/slide{part}.xml'] = xml_bytes(root)

# Retain chart caches, formulas and embedded workbook; change one series color only.
chart_path = 'ppt/slides/charts/chart1.xml'
chart = E.fromstring(parts[chart_path])
series = chart.findall('.//c:ser', NS)
assert series[0].xpath('c:tx//c:v/text()', namespaces=NS) == ['Исходные оценки']
colors = series[0].findall('.//a:srgbClr', NS)
assert [c.get('val') for c in colors] == ['70858A']
colors[0].set('val', 'C95236')
parts[chart_path] = xml_bytes(chart)

backup_notes = {
    21: qa[0]['answer'] + qa[1]['answer'] + qa[2]['answer'] + qa[3]['answer'] + qa[10]['answer'],
    22: qa[12]['answer'],
    23: qa[7]['answer'] + qa[8]['answer'],
    24: qa[9]['answer'] + qa[11]['answer'] + qa[14]['answer'] + qa[16]['answer'],
    25: qa[17]['answer'],
    26: qa[4]['answer'] + qa[5]['answer'] + qa[6]['answer']
}
for number, part in enumerate(ORDER, 1):
    n = f'ppt/notesSlides/notesSlide{part}.xml'
    notes = E.fromstring(parts[n])
    body = next(sp.find('p:txBody', NS) for sp in notes.findall('.//p:sp', NS)
                if sp.find('p:nvSpPr/p:nvPr/p:ph', NS) is not None
                and sp.find('p:nvSpPr/p:nvPr/p:ph', NS).get('type') == 'body')
    old = [''.join(p.xpath('.//a:t/text()', namespaces=NS)) for p in body.findall('a:p', NS)]
    sources = old[old.index('Источники'):] if 'Источники' in old else []
    paragraphs = speech[number] if number <= 20 else backup_notes[number]
    text_body(body, paragraphs + sources)
    parts[n] = xml_bytes(notes)

core = E.fromstring(parts['docProps/core.xml'])
for child in core:
    name = E.QName(child).localname
    if name == 'lastModifiedBy': child.text = None
    if name == 'modified': child.text = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
parts['docProps/core.xml'] = xml_bytes(core)
with zipfile.ZipFile(BUILD / 'candidate.pptx', 'w', zipfile.ZIP_DEFLATED) as z:
    for name, data in parts.items(): z.writestr(name, data)

# Reuse the speech document's page setup and styles, replacing only its body.
source_doc.styles['Normal'].paragraph_format.space_after = Pt(5)
source_doc.styles['Heading 1'].paragraph_format.space_before = Pt(10)
source_doc.styles['Subtitle'].paragraph_format.space_after = Pt(4)
for child in list(source_doc._element.body):
    if E.QName(child).localname != 'sectPr': source_doc._element.body.remove(child)
source_doc.add_paragraph('Текст выступления DAMDID 2026', 'Title')
source_doc.add_paragraph('__REMOVED_IDENTIFIER__\nДоразметка и контролируемые пулы кандидатов', 'Subtitle')
source_doc.add_paragraph('Основной доклад: слайды 1-20. Резервные слайды 21-26 предназначены для ответов на вопросы. Временные отметки дают ориентир для репетиции; паузы для разбора графиков можно увеличить в пределах отведенного времени.')
durations = [5 * math.ceil((sum(len(p.split()) for p in speech[n]) / 125 * 60 + 3) / 5) for n in range(1, 21)]
clock = 0
def mmss(t): return f'{t // 60}:{t % 60:02d}'
for n, duration in enumerate(durations, 1):
    source_doc.add_paragraph(f'Слайд {n}  {titles[n]}', 'Heading 1')
    source_doc.add_paragraph(f'{mmss(clock)}-{mmss(clock + duration)}', 'Subtitle')
    clock += duration
    for p in speech[n]: source_doc.add_paragraph(clean(p))
source_doc.add_page_break()
source_doc.add_paragraph('Ответы на возможные вопросы', 'Title')
source_doc.add_paragraph('Пояснения к данным, модели-судье и сравнению поисковых систем. Указаны слайды, которые можно показать при ответе.')
for item in qa:
    source_doc.add_paragraph(item['q'], 'Heading 2')
    label = 'Резервный слайд' if item['slide'] > 20 else 'Слайд'
    source_doc.add_paragraph(f'{label} {item["slide"]}', 'Subtitle')
    for p in item['answer']: source_doc.add_paragraph(clean(p))
source_doc.core_properties.last_modified_by = ''
source_doc.save(OUT / DOC)

changed = [n for n in parts if parts[n] != original[n]]
assert not any('/embeddings/' in n or '/media/' in n for n in changed)
assert [n for n in changed if '/charts/' in n] == [chart_path]
audit = {'changed_parts': changed, 'slides': 26, 'speech_words': sum(len(p.split()) for v in speech.values() for p in v),
         'reading_duration_seconds': clock, 'qa_count': len(qa), 'speech': speech,
         'qa': qa, 'chart_original_color': '70858A', 'chart_revised_color': 'C95236'}
(BUILD / 'edits.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2))
print(json.dumps({k: audit[k] for k in ('slides', 'speech_words', 'reading_duration_seconds', 'qa_count')}, ensure_ascii=False))
