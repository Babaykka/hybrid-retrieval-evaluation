from pathlib import Path
from copy import deepcopy
from datetime import datetime, timezone
import hashlib
import json
import math
import re
import zipfile

from docx import Document
from lxml import etree as E


ROOT = Path(__file__).resolve().parents[3]
BUILD = ROOT / 'presentation_2026/review_2026-10-12/build/clarity'
SOURCE = ROOT / 'DAMDID_2026_Правки_к_докладу'
OUT = ROOT / 'DAMDID_2026_Уточненная_версия'
PPT = 'DAMDID_2026_Презентация_RU.pptx'
DOC = 'DAMDID_2026_Текст_выступления_RU.docx'
BUILD.mkdir(exist_ok=True)
OUT.mkdir(exist_ok=True)
assert hashlib.sha256((SOURCE / PPT).read_bytes()).hexdigest() == 'aab136cff068a089406bb52d54b7e2a78b01dac8e65e833c180e8b47467307d8'
NS = {'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart'}
ORDER = list(range(1, 11)) + [21, 22, 11, 23, 12, 13, 14, 15, 24, 25, 26, 16, 17, 18, 19, 20]


def clean(text):
    text = text.replace('ё', 'е').replace('Ё', 'Е').replace('—', '-')
    return re.sub(r'(?<=\d) (?=\d{3}(?:\D|$))', '\u00a0', text)


doc = Document(SOURCE / DOC)
speech, titles, qa = {}, {}, []
section, question = None, None
for p in doc.paragraphs:
    if p.style.name == 'Title':
        section, question = None, None
    elif p.style.name == 'Heading 1':
        m = re.match(r'Слайд (\d+)\s+(.+)', p.text)
        section = int(m[1])
        titles[section] = m[2]
        speech[section] = []
    elif p.style.name == 'Heading 2':
        section = None
        question = {'q': p.text, 'answer': [], 'slide': None}
        qa.append(question)
    elif question is not None and p.style.name == 'Subtitle':
        question['slide'] = int(re.search(r'\d+', p.text)[0])
    elif question is not None and p.style.name == 'Normal' and p.text:
        question['answer'].append(p.text)
    elif section is not None and p.style.name == 'Normal' and p.text:
        speech[section].append(p.text)
assert len(speech) == 20 and len(qa) == 19

speech[8][2] = speech[8][2].replace(
    'доверительный интервал показан под графиком.',
    'под графиком приведен 95-процентный доверительный интервал этого прироста.')
speech[15][2] = speech[15][2].replace('разность NDCG: RRF минус E5.',
    'разность двух значений NDCG@10: значение для RRF минус значение для E5.')
speech[15][3] = (
    'Горизонтальные отрезки показывают 95-процентные доверительные интервалы разности. '
    'Здесь одновременно меняются и охват разметки, и входные списки для слияния. '
    'Чтобы выделить влияние охвата, дальше зафиксируем ранжирование.')
speech[16][1] = (
    'Даже при такой неизменной выдаче преимущество переходит от одной системы к другой. '
    'Из NDCG@10 для RRF вычитаем NDCG@10 для E5. Получаем примерно минус 0,066 '
    'при K, равном десяти, и плюс 0,008 при K, равном тридцати. '
    'Отрицательная разность означает преимущество E5, положительная преимущество RRF.')
titles[18] = 'Полнота разметки влияет на сравнение систем'
speech[18] = [
    'Ручная доразметка существенно меняет оценку качества поиска. В наших ручных пулах '
    'NDCG@10 системы E5 увеличивается с 0,488 до 0,879 при неизменной выдаче. '
    'Система уже находила полезные фрагменты, но неполная исходная разметка не давала ей за них баллы.',
    'В масштабном эксперименте при той же выдаче меняется даже победитель сравнения: '
    'на малом пуле E5 опережает RRF, а при расширении пула вперед выходит RRF. '
    'Практический смысл в том, что выбор лучшей системы зависит от того, какие найденные '
    'фрагменты получили оценки. Различие возникает без переобучения поисковых моделей.',
    'Близкие метрики при неполной разметке сами по себе не доказывают одинаковое качество систем. '
    'Но и утверждать, что мы установили истинного лидера, пока нельзя: масштабное сравнение '
    'опирается на оценки одного судьи. Поэтому при сравнении систем нужно проверять, '
    'как вывод зависит от охвата разметки и источника оценок.',
    'Следующий шаг состоит в улучшении автоматической доразметки. Планируем сравнить более '
    'сильные модели-судьи и проверить вариант, который получает запрос, известный релевантный '
    'ответ и фрагмент-кандидат. Качество такой доразметки нужно подтвердить независимыми '
    'человеческими оценками, в том числе на глубоких пулах.'
]
qa.extend([
    {'q': 'Что означает ДИ 95% и учитывает ли он ошибки судьи?', 'slide': 15,
     'answer': [
         'ДИ означает доверительный интервал. На графике показан парный бутстрэп по запросам: '
         'многократно набираем запросы с возвращением и каждый раз считаем среднюю разность '
         'метрик двух систем. Парность означает, что в каждом повторе сравниваем системы '
         'на одних и тех же запросах. Уровень доверия составляет 95%.',
         'В масштабном эксперименте модель-судья и ее предсказания при этом фиксированы. '
         'Интервал показывает вариативность сравнения по запросам, но не включает ошибки '
         'псевдоразметки и неопределенность обучения судьи. Узкий интервал не подтверждает '
         'правильность каждой метки. Для этого нужна независимая человеческая проверка.'
     ]},
    {'q': 'Почему важно, что RRF и E5 меняются местами?', 'slide': 16,
     'answer': [
         'По метрикам часто выбирают систему для дальнейшего использования. Если победитель '
         'меняется только потому, что оценили больше найденных фрагментов, вывод зависит не '
         'только от поисковой системы, но и от покрытия разметкой. В нашем контролируемом '
         'сравнении обе выдачи фиксированы: при K = 10 преимущество у E5, а при K = 30 у RRF.',
         'Это показывает чувствительность сравнения к пулу оценок. Само по себе такое '
         'переставление не доказывает, что RRF объективно лучше, поскольку оценки дает один '
         'судья. Также отсюда нельзя сделать вывод о существенных различиях всех современных '
         'SOTA-моделей. Проверяемый вывод уже: близкие показатели на неполной разметке не '
         'достаточны, чтобы считать системы равноценными. Для надежного выбора нужна более '
         'полная и независимо проверенная разметка.'
     ]}
])

with zipfile.ZipFile(SOURCE / PPT) as z:
    parts = {n: z.read(n) for n in z.namelist()}
original = dict(parts)


def xml_bytes(root):
    return E.tostring(root, encoding='UTF-8', xml_declaration=True, standalone=True)


def text_body(body, lines):
    old = body.findall('a:p', NS)
    template = old[0] if old else None
    for p in old:
        body.remove(p)
    for line in lines:
        p = E.SubElement(body, f'{{{NS["a"]}}}p')
        if template is not None and template.find('a:pPr', NS) is not None:
            p.append(deepcopy(template.find('a:pPr', NS)))
        r = E.SubElement(p, f'{{{NS["a"]}}}r')
        if template is not None and template.find('a:r/a:rPr', NS) is not None:
            r.append(deepcopy(template.find('a:r/a:rPr', NS)))
        E.SubElement(r, f'{{{NS["a"]}}}t').text = clean(line)


def shape(root, shape_id):
    return next(s for s in root.findall('.//p:sp', NS)
                if s.find('p:nvSpPr/p:cNvPr', NS).get('id') == str(shape_id))


def set_text(root, shape_id, text):
    text_body(shape(root, shape_id).find('p:txBody', NS), text.split('\n'))


def formula(body):
    # Preserve native rich text and use true subscripts for the system names.
    p = body.find('a:p', NS)
    style = deepcopy(p.find('a:r/a:rPr', NS))
    if style is None:
        style = deepcopy(p.find('a:pPr/a:defRPr', NS))
        style.tag = f'{{{NS["a"]}}}rPr'
    for child in list(p):
        if child.tag != f'{{{NS["a"]}}}pPr':
            p.remove(child)
    for text, subscript in [('NDCG@10', False), ('RRF', True),
                            (' − NDCG@10', False), ('E5', True)]:
        r = E.SubElement(p, f'{{{NS["a"]}}}r')
        rp = deepcopy(style)
        if subscript:
            rp.set('baseline', '-25000')
        r.append(rp)
        E.SubElement(r, f'{{{NS["a"]}}}t').text = text


roots = {n: E.fromstring(parts[f'ppt/slides/slide{n}.xml']) for n in [8, 9, 10, 21, 12, 13, 15, 16]}
set_text(roots[8], 5, 'Изменение E5: +0,391   ДИ 95% [0,365; 0,417]')
style = deepcopy(shape(roots[9], 12).find('.//a:rPr', NS))
sp = shape(roots[9], 12)
sp.getparent().remove(sp)
for rp in shape(roots[10], 10).findall('.//a:rPr', NS):
    rp.getparent().replace(rp, deepcopy(style))
for color in shape(roots[10], 9).findall('.//a:rPr/a:solidFill/a:srgbClr', NS):
    color.set('val', '007F78')
sp = shape(roots[21], 8)
sp.getparent().remove(sp)
set_text(roots[12], 7, 'Один судья. Парный ДИ 95% [0,0082; 0,0095].')
formula(shape(roots[13], 3).find('p:txBody', NS))
set_text(roots[15], 12, titles[18])
set_text(roots[15], 3, 'Ручная доразметка меняет оценку качества')
set_text(roots[15], 7, 'При фиксированной выдаче и одном судье')
set_text(roots[15], 8, 'Вывод о лучшей системе зависит от охвата разметки')
set_text(roots[15], 10, 'Направления дальнейшего исследования\nУлучшить судью и проверить доразметку по оценкам человека')
for node in roots[16].findall('.//a:t', NS):
    if node.text == '95%-ный парный ДИ':
        node.text = 'Парный ДИ 95%'
for n, root in roots.items():
    parts[f'ppt/slides/slide{n}.xml'] = xml_bytes(root)

chart_path = 'ppt/slides/charts/chart4.xml'
chart = E.fromstring(parts[chart_path])
axis_body = next(b for b in chart.findall('.//c:title/c:tx/c:rich', NS)
                 if 'RRF минус E5' in ''.join(b.xpath('.//a:t/text()', namespaces=NS)))
formula(axis_body)
parts[chart_path] = xml_bytes(chart)

for n in [8, 15, 16, 18]:
    part = ORDER[n - 1]
    name = f'ppt/notesSlides/notesSlide{part}.xml'
    notes = E.fromstring(parts[name])
    body = next(s.find('p:txBody', NS) for s in notes.findall('.//p:sp', NS)
                if s.find('p:nvSpPr/p:nvPr/p:ph', NS) is not None
                and s.find('p:nvSpPr/p:nvPr/p:ph', NS).get('type') == 'body')
    old = [''.join(p.xpath('.//a:t/text()', namespaces=NS)) for p in body.findall('a:p', NS)]
    sources = old[old.index('Источники'):] if 'Источники' in old else []
    text_body(body, speech[n] + sources)
    parts[name] = xml_bytes(notes)

core = E.fromstring(parts['docProps/core.xml'])
for child in core:
    if E.QName(child).localname == 'lastModifiedBy': child.text = None
    if E.QName(child).localname == 'modified':
        child.text = datetime.now(timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')
parts['docProps/core.xml'] = xml_bytes(core)
with zipfile.ZipFile(BUILD / 'candidate.pptx', 'w', zipfile.ZIP_DEFLATED) as z:
    for name, data in parts.items():
        z.writestr(name, data)

for child in list(doc._element.body):
    if E.QName(child).localname != 'sectPr': doc._element.body.remove(child)
doc.add_paragraph('Текст выступления DAMDID 2026', 'Title')
doc.add_paragraph('__REMOVED_IDENTIFIER__\nДоразметка и контролируемые пулы кандидатов', 'Subtitle')
doc.add_paragraph('Основной доклад: слайды 1-20. Резервные слайды 21-26 предназначены для ответов на вопросы. Временные отметки дают ориентир для репетиции; паузы для разбора графиков можно увеличить в пределах отведенного времени.')
durations = [5 * math.ceil((sum(len(p.split()) for p in speech[n]) / 125 * 60 + 3) / 5) for n in range(1, 21)]
clock = 0
def mmss(t): return f'{t // 60}:{t % 60:02d}'
for n, duration in enumerate(durations, 1):
    doc.add_paragraph(f'Слайд {n}  {titles[n]}', 'Heading 1')
    doc.add_paragraph(f'{mmss(clock)}-{mmss(clock + duration)}', 'Subtitle')
    clock += duration
    for p in speech[n]: doc.add_paragraph(clean(p))
doc.add_page_break()
doc.add_paragraph('Ответы на возможные вопросы', 'Title')
doc.add_paragraph('Пояснения к данным, модели-судье и сравнению поисковых систем. Указаны слайды, которые можно показать при ответе.')
for item in qa:
    doc.add_paragraph(item['q'], 'Heading 2')
    label = 'Резервный слайд' if item['slide'] > 20 else 'Слайд'
    doc.add_paragraph(f'{label} {item["slide"]}', 'Subtitle')
    for p in item['answer']: doc.add_paragraph(clean(p))
doc.core_properties.last_modified_by = ''
doc.save(OUT / DOC)

changed = [n for n in parts if parts[n] != original[n]]
assert not any('/embeddings/' in n or '/media/' in n for n in changed)
audit = {'changed_parts': changed, 'speech': speech, 'qa': qa, 'titles': titles,
         'reading_duration_seconds': clock, 'qa_count': len(qa)}
(BUILD / 'edits.json').write_text(json.dumps(audit, ensure_ascii=False, indent=2))
print(json.dumps({'changed_parts': changed, 'qa_count': len(qa), 'reading_seconds': clock}, ensure_ascii=False))
