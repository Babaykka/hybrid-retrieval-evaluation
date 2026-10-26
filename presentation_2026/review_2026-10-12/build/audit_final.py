import os
from pathlib import Path
from lxml import etree as E
import hashlib
import json
import io
import re
import zipfile
from docx import Document

PRIVATE_IDENTITY_TOKENS = tuple(token.strip() for token in os.environ.get("PRIVATE_IDENTITY_TOKENS", "").split("|") if token.strip())

B = Path(__file__).parent
F = B / 'validated/validated_v3.pptx'
O = Path('./DAMDID_2026_После_рецензии')
NS = {'p':'http://schemas.openxmlformats.org/presentationml/2006/main', 'a':'http://schemas.openxmlformats.org/drawingml/2006/main'}
with zipfile.ZipFile(B / 'source/presentation.pptx') as src, zipfile.ZipFile(F) as dst:
    assert src.testzip() is None and dst.testzip() is None
    # Source numerical evidence must survive the structural edits byte-for-byte.
    protected = [n for n in src.namelist() if '/charts/' in n or n.startswith(('ppt/embeddings/','ppt/media/'))]
    assert all(src.read(n) == dst.read(n) for n in protected)
    table_count = 0
    for i in range(1,21):
        before = E.fromstring(src.read(f'ppt/slides/slide{i}.xml')).findall('.//a:tbl',NS)
        after = E.fromstring(dst.read(f'ppt/slides/slide{i}.xml')).findall('.//a:tbl',NS)
        assert len(before) == len(after)
        for x,y in zip(before,after):
            a = [''.join(t.itertext()) for t in x.findall('.//a:t',NS)]
            b = [''.join(t.itertext()) for t in y.findall('.//a:t',NS)]
            assert a == b, (i,a,b)
            table_count += 1
    tokens = []
    for n in dst.namelist():
        if n.endswith('.xml'):
            tree = E.fromstring(dst.read(n))
            tokens.extend(tree.itertext())
    combined = '\n'.join(tokens)
    for n in dst.namelist():
        if n.endswith('.xlsx'):
            with zipfile.ZipFile(io.BytesIO(dst.read(n))) as workbook:
                for entry in workbook.namelist():
                    if entry.endswith('.xml'):
                        combined += '\n' + '\n'.join(E.fromstring(workbook.read(entry)).itertext())
    for word in ['Грач Бабаян', 'ChatGPT', 'TODO', 'Study results:', 'revision_2026/', 'presentation_2026/', '/Users/', *PRIVATE_IDENTITY_TOKENS]:
        assert word.casefold() not in combined.casefold(), word
    assert len([n for n in dst.namelist() if re.fullmatch(r'ppt/slides/slide\d+\.xml',n)]) == 26
    for i in range(1,27):
        assert f'ppt/notesSlides/notesSlide{i}.xml' in dst.namelist()
    d = Document(O / 'DAMDID_2026_Текст_выступления_RU.docx')
    headings = [p.text for p in d.paragraphs if p.style.name == 'Heading 1']
    assert len(headings) == 20
    assert [int(re.match(r'Слайд (\d+)',h)[1]) for h in headings] == list(range(1,21))
    manifest = json.loads((B/'slide_manifest.json').read_text())
    for item in manifest[:20]:
        n = item['source_part']
        text = '\n'.join(E.fromstring(dst.read(f'ppt/notesSlides/notesSlide{n}.xml')).itertext())
        for para in item['speech']: assert para in text
    result = {'status':'pass','slides':26,'protected_chart_workbook_media_parts':len(protected),
              'original_tables_preserved':table_count,'main_speech_sections':len(headings),
              'questions':len([p for p in d.paragraphs if p.style.name=='Heading 2']),
              'pptx_sha256':hashlib.sha256(F.read_bytes()).hexdigest()}
    (B/'audit_final.json').write_text(json.dumps(result,indent=2))
    print(json.dumps(result,indent=2))
