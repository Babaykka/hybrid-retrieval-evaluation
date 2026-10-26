"""Verify final package consistency and leave only client files in the ZIP."""
import os
import hashlib
import io
import json
import posixpath
import re
from pathlib import Path
from zipfile import ZipFile, ZIP_DEFLATED

from docx import Document
from lxml import etree as E
from pypdf import PdfReader

PRIVATE_IDENTITY_TOKENS = tuple(token.strip() for token in os.environ.get("PRIVATE_IDENTITY_TOKENS", "").split("|") if token.strip())

ROOT = Path(__file__).resolve().parents[2]
B = Path(__file__).parent
OUT = ROOT / 'DAMDID_2026_Проверено_13-10'
SOURCE = ROOT / 'DAMDID_2026_Уточненная_версия 2'
PPT = 'DAMDID_2026_Презентация_RU.pptx'
DOC = 'DAMDID_2026_Текст_выступления_RU.docx'
PDF = 'DAMDID_2026_Презентация_RU.pdf'
NS = {'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
      'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
      'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
      'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'}
bad_name = re.compile(r'__REMOVED_IDENTIFIER__|hratch|__REMOVED_IDENTIFIER__|грач|грах|бабаян', re.I)


def sha(path):
    return hashlib.sha256(path.read_bytes()).hexdigest()


def norm(s):
    return ' '.join(s.split())


def rels(z, part):
    d, n = posixpath.split(part)
    return {r.get('Id'): (r.get('Target').lstrip('/') if r.get('Target').startswith('/') else posixpath.normpath(d + '/' + r.get('Target'))) for r in E.fromstring(z.read(d + '/_rels/' + n + '.rels'))}


def scan(data, path):
    if data.startswith(b'PK'):
        with ZipFile(io.BytesIO(data)) as z:
            assert z.testzip() is None
            for n in z.namelist():
                scan(z.read(n), path + '::' + n)
    elif path.endswith(('.xml', '.rels')):
        assert not bad_name.search(data.decode('utf-8')), path


for name, expected in json.loads((B / 'edits.json').read_text())['source_sha256'].items():
    assert sha(SOURCE / name) == expected, 'Source was modified: ' + name

old_doc = Document(SOURCE / DOC)
doc = Document(OUT / DOC)
old_text = '\n'.join(p.text for p in old_doc.paragraphs)
new_text = '\n'.join(p.text for p in doc.paragraphs)
assert re.findall(r'\d+', old_text) == re.findall(r'\d+', new_text), 'A speech number changed'
assert 'не достаточны' not in new_text
assert 'верных отрицательных решения' not in new_text
assert 'ё' not in new_text and '—' not in new_text
assert all((token not in str(new_text) for token in PRIVATE_IDENTITY_TOKENS))

speeches = {}
section = None
for p in doc.paragraphs:
    if p.style.name == 'Heading 1':
        section = int(re.match(r'Слайд (\d+)', p.text)[1])
        speeches[section] = []
    elif p.style.name == 'Title':
        section = None
    elif section and p.style.name == 'Normal' and p.text:
        speeches[section].append(p.text)

with ZipFile(OUT / PPT) as z, ZipFile(SOURCE / PPT) as old:
    for n in z.namelist():
        if '/embeddings/' in n or '/media/' in n:
            assert z.read(n) == old.read(n), 'Original media/data changed: ' + n
    prels = rels(z, 'ppt/presentation.xml')
    parts = [prels[s.get('{' + NS['r'] + '}id')] for s in E.fromstring(z.read('ppt/presentation.xml')).findall('p:sldIdLst/p:sldId', NS)]
    for i, part in enumerate(parts, 1):
        r = E.fromstring(z.read(part))
        text = ''.join(r.xpath('//a:t/text()', namespaces=NS))
        before = ''.join(E.fromstring(old.read(part)).xpath('//a:t/text()', namespaces=NS))
        assert text == before.replace('ё', 'е').replace('Ё', 'Е'), i
        assert not any(s in text for s in ['TODO', 'ChatGPT', 'queries.dev.tsv', '0.597', '0.967'])
        if i <= 20:
            notes_path = next(v for v in rels(z, part).values() if '/notesSlides/' in v)
            notes = E.fromstring(z.read(notes_path))
            body = next(sp.find('p:txBody', NS) for sp in notes.findall('.//p:sp', NS) if sp.find('p:nvSpPr/p:nvPr/p:ph', NS) is not None and sp.find('p:nvSpPr/p:nvPr/p:ph', NS).get('type') == 'body')
            lines = [''.join(p.xpath('.//a:t/text()', namespaces=NS)) for p in body.findall('a:p', NS)]
            if 'Источники' in lines:
                lines = lines[:lines.index('Источники')]
            assert list(map(norm, lines)) == list(map(norm, speeches[i])), 'Notes mismatch: ' + str(i)

reader = PdfReader(OUT / PDF)
assert len(reader.pages) == 26
assert len(PdfReader(B / 'docx_render' / (Path(DOC).stem + '.pdf')).pages) == 12
pdf_text = '\n'.join(p.extract_text() for p in reader.pages)
assert not bad_name.search(pdf_text + str(reader.metadata))
for i, values in {8: ['0,248', '0,488', '0,644', '0,879'], 10: ['0,907', '0,875', '0,2187'], 15: ['0,0088', '0,0082', '0,0095'], 16: ['14,4%', '145 643', '0,0661', '0,0080'], 20: ['__REMOVED_IDENTIFIER__']}.items():
    page = norm(reader.pages[i - 1].extract_text())
    assert all(v in page for v in values), (i, page)
for name in [PPT, DOC]:
    scan((OUT / name).read_bytes(), name)

archive = OUT.with_suffix('.zip')
with ZipFile(archive, 'x', ZIP_DEFLATED) as z:
    for name in [PPT, PDF, DOC]:
        z.write(OUT / name, name)
with ZipFile(archive) as z:
    assert z.testzip() is None
    assert len(z.namelist()) == 3
    for name in z.namelist():
        assert z.read(name) == (OUT / name).read_bytes()

report = {'status': 'pass', 'slides': 26, 'speech_pages': 12, 'main_notes_match_speech': 20,
          'source_preserved': True, 'numeric_audit': json.loads((B / 'numbers_verified.json').read_text())['status'],
          'metadata_scan': 'pass', 'archive_integrity': 'pass',
          'sha256': {p.name: sha(p) for p in [OUT / PPT, OUT / PDF, OUT / DOC, archive]}}
(B / 'delivery_verified.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps(report, ensure_ascii=False, indent=2))
