"""Check the delivered scientific values against saved per-query results."""
import io
import json
import posixpath
import re
import sys
from pathlib import Path
from zipfile import ZipFile

import numpy as np
from lxml import etree as E
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

ROOT = Path(__file__).resolve().parents[2]
RESULTS = ROOT / 'revision_2026/results'
DECK = Path(sys.argv[1])
NS = {
    'c': 'http://schemas.openxmlformats.org/drawingml/2006/chart',
    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
    'p': 'http://schemas.openxmlformats.org/presentationml/2006/main',
    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
}
m = json.loads((RESULTS / 'manual_summary.json').read_text())
f = json.loads((RESULTS / 'fullset_summary.json').read_text())
roc = json.loads((ROOT / 'presentation_2026/premium/reviews/test_roc.json').read_text())
checks = []


def check(name, actual, expected, atol=1e-11):
    np.testing.assert_allclose(actual, expected, rtol=0, atol=atol, err_msg=name)
    checks.append(name)


def resolve(part, target):
    return target.lstrip('/') if target.startswith('/') else posixpath.normpath(posixpath.join(posixpath.dirname(part), target))


def relations(z, part):
    d, name = posixpath.split(part)
    return {r.get('Id'): resolve(part, r.get('Target')) for r in E.fromstring(z.read(d + '/_rels/' + name + '.rels'))}


def nums(series, kind):
    return [float(x) for x in series.xpath(f'./c:{kind}/c:numRef/c:numCache/c:pt/c:v/text() | ./c:{kind}/c:numLit/c:pt/c:v/text()', namespaces=NS)]


def ru(value, digits):
    return f'{value:.{digits}f}'.replace('.', ',')


with np.load(RESULTS / 'manual_per_query.npz') as data:
    check('manual query count', len(data['qids']), m['queries'])
    check('manual clean count', data['clean'].sum(), m['clean_queries'])
    for j, model in enumerate(data['models']):
        for labels in ['official', 'human', 'judge']:
            check(f'manual {model} {labels}', data[labels][:, j].mean(), m['metrics'][str(model)][labels])
        check(f'manual delta {model}', (data['human'][:, j] - data['official'][:, j]).mean(), m['metrics'][str(model)]['human_minus_official']['mean'])
        check(f'manual clean {model}', data['human'][data['clean'], j].mean(), m['metrics'][str(model)]['clean_queries_human'])

with np.load(RESULTS / 'fullset_perquery.npz') as data:
    check('full query count', len(np.unique(data['query_ids'])), 101093)
    masks = {'all_queries': np.ones(len(data['query_ids']), dtype=bool), 'excluding_manual_queries': ~data['manual_overlap'], 'manual_overlap_only': data['manual_overlap']}
    check('query subtraction', 101093 - 1179, masks['excluding_manual_queries'].sum())
    for population, mask in masks.items():
        for i, depth in enumerate(f['populations'][population]['depths']):
            check(f'{population} query count', mask.sum(), depth['queries'])
            for arr, total, mean in [('candidate_count', 'candidate_pairs', 'candidates_mean_all_queries'), ('positive_count', 'positive_pairs', 'positives_mean_all_queries')]:
                check(f'{population} K={depth["k"]} {total}', data[arr][mask, i].sum(), depth[total])
                check(f'{population} K={depth["k"]} {mean}', data[arr][mask, i].mean(), depth[mean])
            for j, model in enumerate(data['model_names']):
                v = depth['models'][str(model)]
                check(f'{population} K={depth["k"]} {model}', data['ndcg'][mask, i, j].mean(), v['ndcg10_mean_all_queries'])
                check(f'{population} K={depth["k"]} outside {model}', data['top10_outside_pool'][mask, i, j].sum(), v['top10_outside_pool_total'])
    e = list(data['model_names']).index('e5')
    for population, mask in masks.items():
        for r in f['paired_inference'][population]:
            model = r['contrast'].removesuffix('_minus_e5')
            j = list(data['model_names']).index(model)
            k = list(data['ks']).index(r['k'])
            check(f'{population} paired {r["contrast"]} K={r["k"]}', (data['ndcg'][mask, k, j] - data['ndcg'][mask, k, e]).mean(), r['mean_delta'])

with np.load(RESULTS / 'fullset_inference_replicates.npz') as data:
    for population in masks:
        samples = data[population + '_bootstrap_mean_delta']
        for i, (k, model) in enumerate(zip(data['contrast_k'], data['contrast_model'])):
            row = next(r for r in f['paired_inference'][population] if r['k'] == k and r['contrast'] == str(model) + '_minus_e5')
            check(f'{population} CI {model} K={k}', np.quantile(samples[:, i], [.025, .975]), row['bootstrap_ci95'])

test = m['judge_rescored_split_metrics']['test']
tn, fp, fn, tp = np.asarray(test['confusion_matrix']).ravel()
check('test matrix total', tn + fp + fn + tp, test['n'])
check('test precision', tp / (tp + fp), test['precision'])
check('test recall', tp / (tp + fn), test['recall'])
check('test F1', 2 * tp / (2 * tp + fp + fn), test['f1'])
check('judge query split', sum(s['queries'] for s in m['judge_rescored_split_metrics'].values()), 1179)
check('judge pair split', sum(s['n'] for s in m['judge_rescored_split_metrics'].values()), 33667)
check('manual candidates per query', 33667 / 1179, m['pool_mean'])
for labels in ['official', 'human']:
    check('positive density ' + labels, m['density'][labels]['pairs'] / 1179, m['density'][labels]['mean'])
check('invalid label pairs', len(m['invalid_labels']), 29)
check('invalid label queries', len({r[0] for r in m['invalid_labels']}), 18)
check('empty text queries', len({r[0] for r in m['empty_texts']}), 5)
check('union of excluded queries', len({r[0] for r in m['invalid_labels'] + m['empty_texts']}), 23)
check('remaining manual queries', 1179 - 23, 1156)
check('leaderboard spread', .450 - .436, .014)
check('RRF common example', round(3 / 71, 4), .0423)
check('RRF single example', round(1 / 61, 4), .0164)

depths = f['populations']['all_queries']['depths']
out = depths[0]['models']['rrf_fixed200']
check('outside pool count', out['top10_outside_pool_total'], 145643)
check('outside pool percentage', round(100 * out['top10_outside_pool_total'] / out['top10_count_total'], 1), 14.4)

with ZipFile(DECK) as z:
    assert z.testzip() is None
    rels = relations(z, 'ppt/presentation.xml')
    slides = [rels[n.get('{' + NS['r'] + '}id')] for n in E.fromstring(z.read('ppt/presentation.xml')).findall('p:sldIdLst/p:sldId', NS)]
    assert len(slides) == 26
    charts = []
    for number in [8, 10, 13, 15, 16]:
        part = slides[number - 1]
        rid = E.fromstring(z.read(part)).find('.//c:chart', NS).get('{' + NS['r'] + '}id')
        chart_part = relations(z, part)[rid]
        chart = E.fromstring(z.read(chart_part))
        charts.append(chart)
        rid = chart.find('c:externalData', NS).get('{' + NS['r'] + '}id')
        wb = load_workbook(io.BytesIO(z.read(relations(z, chart_part)[rid])), data_only=True)
        for ref in chart.findall('.//c:numRef', NS):
            formula = ref.findtext('c:f', namespaces=NS)
            sheet, cell_range = formula.rsplit('!', 1)
            sheet = sheet.strip("'").replace("''", "'")
            c1, r1, c2, r2 = range_boundaries(cell_range.replace('$', ''))
            cells = [wb[sheet].cell(r, c).value for r in range(r1, r2 + 1) for c in range(c1, c2 + 1)]
            pts = ref.findall('c:numCache/c:pt', NS)
            for pt in pts:
                check(f'chart {number} workbook {formula} {pt.get("idx")}', float(pt.findtext('c:v', namespaces=NS)), float(cells[int(pt.get('idx'))]), 1e-8)
        wb.close()
    for series, label in zip(charts[0].findall('.//c:ser', NS), ['official', 'human']):
        check('manual chart ' + label, nums(series, 'val'), [m['metrics'][k][label] for k in ['rrf3', 'bge', 'e5', 'bm25']], 1e-8)
    chance, curve, point = charts[1].findall('.//c:ser', NS)
    for axis, key in [('xVal', 'fpr'), ('yVal', 'tpr')]:
        check('ROC ' + axis, nums(curve, axis), [p[key] for p in roc['points']], 1e-8)
        check('ROC threshold ' + axis, nums(point, axis), [roc['operating_point'][key]] * 80, 1e-8)
    for series, key in zip(charts[2].findall('.//c:ser', NS), ['candidates_mean_all_queries', 'positives_mean_all_queries']):
        check('pool chart ' + key, nums(series, 'yVal'), [r[key] for r in depths], 1e-8)
    for chart, contrast, horizontal in zip(charts[3:], ['rrf_variable_minus_e5', 'rrf_fixed200_minus_e5'], [True, False]):
        rows = [r for r in f['paired_inference']['all_queries'] if r['contrast'] == contrast]
        series = chart.findall('.//c:ser', NS)
        values = nums(series[-1], 'xVal' if horizontal else 'yVal')
        check('RRF chart ' + contrast, values, [r['mean_delta'] for r in rows], 1e-8)
        errors = series[-1].find('c:errBars', NS)
        for side, sign, bound in [('plus', 1, 1), ('minus', -1, 0)]:
            offsets = [float(x) for x in errors.xpath(f'./c:{side}/c:numLit/c:pt/c:v/text() | ./c:{side}/c:numRef/c:numCache/c:pt/c:v/text()', namespaces=NS)]
            check('chart CI ' + contrast + side, np.array(values) + sign * np.array(offsets), [r['bootstrap_ci95'][bound] for r in rows], 1e-8)
        if horizontal:
            for i, interval in enumerate(series[1:6]):
                lo, hi = rows[i]['bootstrap_ci95']
                check('horizontal CI ' + str(i), nums(interval, 'xVal'), [lo, hi, hi, hi, hi], 1e-8)

    def table(number):
        node = E.fromstring(z.read(slides[number - 1])).find('.//a:tbl', NS)
        return [[''.join(c.xpath('.//a:t/text()', namespaces=NS)) for c in row.findall('a:tc', NS)] for row in node.findall('a:tr', NS)]

    for row, model in zip(table(22)[1:], ['bm25', 'e5', 'bge', 'rrf3']):
        d = m['metrics'][model]['human_minus_official']
        assert row[1] == '+' + ru(d['mean'], 3)
        assert row[2] == ' ... '.join(ru(v, 3) for v in d['ci95'])
    for row, r in zip(table(24)[1:], depths):
        assert row == [str(r['k']), ru(r['candidates_mean_all_queries'], 1), ru(r['positives_mean_all_queries'], 1)] + [ru(r['models'][k]['ndcg10_mean_all_queries'], 3) for k in ['e5', 'rrf_variable', 'rrf_fixed200']], row
    for row, model in zip(table(25)[1:], ['bm25', 'e5', 'bge', 'rrf3']):
        assert row[1:] == [ru(m['metrics'][model][k], 3) for k in ['human', 'clean_queries_human']]
    assert [[int(v.replace(' ', '')) for v in row[1:]] for row in table(23)[1:]] == test['confusion_matrix']

report = {'status': 'pass', 'checks': len(checks), 'scope': 'Saved per-query metrics, bootstrap replicates, chart caches and embedded workbooks, displayed tables, derived arithmetic', 'deck': str(DECK), 'check_names': checks}
Path(__file__).with_name('numbers_verified.json').write_text(json.dumps(report, ensure_ascii=False, indent=2))
print(json.dumps({k: v for k, v in report.items() if k != 'check_names'}, ensure_ascii=False, indent=2))
